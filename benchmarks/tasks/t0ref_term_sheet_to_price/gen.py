"""Generator for t0ref_term_sheet_to_price.

Atomic formula task with a markdown reference file. The Builder must read
`inputs/term_sheet.md` to extract numeric inputs (pre-money, round size,
pre-round FD shares) and then write formulas for post-money, share price,
new investor shares, post-round FD, and Starlight's post-round ownership.

Artifacts produced:
  inputs/term_sheet.md   — authored by hand (not overwritten here)
  stubs/starter.xlsx     — mostly empty Calc sheet (labels only)
  gold/model.xlsx        — stub + populated literals + formulas
  gold/results.json      — semantic key -> Sheet!Cell map
"""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

from recalc import recalc_xlsx

TASK_DIR = Path(__file__).resolve().parent

# --- parameters extracted from term_sheet.md --------------------------------
PRE_MONEY = 50_000_000
ROUND_SIZE = 10_000_000
PRE_FD = 5_000_000

# Expected derived values (for reference; grading.yaml hardcodes these)
POST_MONEY = PRE_MONEY + ROUND_SIZE            # 60M
SHARE_PRICE = PRE_MONEY / PRE_FD                # 10.0
NEW_INVESTOR_SHARES = ROUND_SIZE / SHARE_PRICE  # 1_000_000
POST_FD = PRE_FD + NEW_INVESTOR_SHARES          # 6_000_000
STARLIGHT_PCT_POST = NEW_INVESTOR_SHARES / POST_FD  # 0.16666...


# --- builders ---------------------------------------------------------------


def build_calc_sheet(ws, *, fill: bool) -> None:
    """Lay out the Calc sheet. If fill=True, populate literals + formulas."""
    ws.title = "Calc"
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 18

    ws.cell(row=1, column=1, value="Cobalt Series A — share price").font = Font(bold=True, size=13)

    # Input labels (B3:B5 filled only in gold)
    ws.cell(row=3, column=1, value="Pre-money valuation")
    ws.cell(row=4, column=1, value="New money (round size)")
    ws.cell(row=5, column=1, value="Pre-round FD shares")

    # Derived labels (B8:B12 filled only in gold)
    ws.cell(row=8, column=1, value="Post-money valuation")
    ws.cell(row=9, column=1, value="Share price")
    ws.cell(row=10, column=1, value="New investor shares")
    ws.cell(row=11, column=1, value="Post-round FD")
    ws.cell(row=12, column=1, value="Starlight % post")

    # Number formats (applied whether or not the cell is filled)
    ws.cell(row=3, column=2).number_format = "$#,##0"
    ws.cell(row=4, column=2).number_format = "$#,##0"
    ws.cell(row=5, column=2).number_format = "#,##0"
    ws.cell(row=8, column=2).number_format = "$#,##0"
    ws.cell(row=9, column=2).number_format = "$#,##0.00"
    ws.cell(row=10, column=2).number_format = "#,##0"
    ws.cell(row=11, column=2).number_format = "#,##0"
    ws.cell(row=12, column=2).number_format = "0.00%"

    if fill:
        ws.cell(row=3, column=2, value=PRE_MONEY)
        ws.cell(row=4, column=2, value=ROUND_SIZE)
        ws.cell(row=5, column=2, value=PRE_FD)
        ws.cell(row=8, column=2, value="=B3+B4")
        ws.cell(row=9, column=2, value="=B3/B5")
        ws.cell(row=10, column=2, value="=B4/B9")
        ws.cell(row=11, column=2, value="=B5+B10")
        ws.cell(row=12, column=2, value="=B10/B11")
        # Re-apply number formats since setting .value wipes them on some versions
        ws.cell(row=3, column=2).number_format = "$#,##0"
        ws.cell(row=4, column=2).number_format = "$#,##0"
        ws.cell(row=5, column=2).number_format = "#,##0"
        ws.cell(row=8, column=2).number_format = "$#,##0"
        ws.cell(row=9, column=2).number_format = "$#,##0.00"
        ws.cell(row=10, column=2).number_format = "#,##0"
        ws.cell(row=11, column=2).number_format = "#,##0"
        ws.cell(row=12, column=2).number_format = "0.00%"


def build_stub(path: Path) -> None:
    wb = openpyxl.Workbook()
    build_calc_sheet(wb.active, fill=False)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def build_gold(path: Path) -> None:
    wb = openpyxl.Workbook()
    build_calc_sheet(wb.active, fill=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    recalc_xlsx(path)


def build_results_json(path: Path) -> None:
    data = {
        "pre_money": "Calc!B3",
        "round_size": "Calc!B4",
        "pre_fd": "Calc!B5",
        "post_money": "Calc!B8",
        "share_price": "Calc!B9",
        "new_investor_shares": "Calc!B10",
        "post_fd": "Calc!B11",
        "starlight_pct_post": "Calc!B12",
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, indent=2))


def main() -> None:
    build_stub(TASK_DIR / "stubs" / "starter.xlsx")
    build_gold(TASK_DIR / "gold" / "model.xlsx")
    build_results_json(TASK_DIR / "gold" / "results.json")
    print("t0ref_term_sheet_to_price artifacts written.")
    print(f"  post_money          = {POST_MONEY:,.2f}")
    print(f"  share_price         = {SHARE_PRICE:,.6f}")
    print(f"  new_investor_shares = {NEW_INVESTOR_SHARES:,.2f}")
    print(f"  post_fd             = {POST_FD:,.2f}")
    print(f"  starlight_pct_post  = {STARLIGHT_PCT_POST:,.6f}")


if __name__ == "__main__":
    main()
