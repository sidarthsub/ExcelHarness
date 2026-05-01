"""Generator for t0ref_cap_table_fd_count.

Deterministically produces:
  inputs/messy_cap_table.xlsx  — reference cap table (9 rows, mixed classes)
  stubs/starter.xlsx           — Calc sheet with title + reference-file pointer;
                                 B5 (answer) left blank
  gold/model.xlsx              — Calc sheet filled with =SUM('Cap Table'!C2:C10),
                                 plus a Cap Table sheet mirroring the reference
  gold/results.json            — {"total_fd": "Calc!B5"}

Math:
  Total FD = sum of all 9 rows of the reference Cap Table shares column
           = 12,000,000
"""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from recalc import recalc_xlsx

TASK_DIR = Path(__file__).resolve().parent

# --- data -------------------------------------------------------------------
CAP_ROWS = [
    ("Founder A", "Common", 3_000_000),
    ("Founder B", "Common", 2_500_000),
    ("Seed Lead", "Preferred Seed", 1_500_000),
    ("Seed Angels", "Preferred Seed", 750_000),
    ("Series A Lead", "Preferred A", 2_000_000),
    ("Series A Other", "Preferred A", 1_000_000),
    ("Options granted", "Options", 600_000),
    ("Options unallocated", "Options", 400_000),
    ("Warrants", "Warrants", 250_000),
]

GOLD_TOTAL_FD = sum(r[2] for r in CAP_ROWS)  # 12_000_000


# --- builders ---------------------------------------------------------------


def _write_cap_table(ws) -> None:
    """Write the 9-row cap table (headers row 1, data rows 2..10)."""
    ws.title = "Cap Table"
    hdr_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="E8E8E8")

    headers = ["Holder", "Class", "Shares"]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = hdr_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center")

    for r, (holder, cls, shares) in enumerate(CAP_ROWS, 2):
        ws.cell(row=r, column=1, value=holder)
        ws.cell(row=r, column=2, value=cls)
        ws.cell(row=r, column=3, value=shares).number_format = "#,##0"

    for col, width in [("A", 26), ("B", 22), ("C", 14)]:
        ws.column_dimensions[col].width = width


def build_reference(path: Path) -> None:
    wb = openpyxl.Workbook()
    _write_cap_table(wb.active)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    recalc_xlsx(path)


def build_stub(path: Path) -> None:
    """Starter workbook: Calc sheet with title, reference pointer, blank B5."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Calc"

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 28

    ws.cell(row=1, column=1, value="Cap Table FD count").font = Font(bold=True, size=13)

    ws.cell(row=3, column=1, value="Reference file")
    ws.cell(row=3, column=2, value="inputs/messy_cap_table.xlsx")

    ws.cell(row=5, column=1, value="Total FD (from reference)").font = Font(bold=True)
    # B5 intentionally blank — the Builder must fill it with a formula.
    ws.cell(row=5, column=2).number_format = "#,##0"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def build_gold(path: Path) -> None:
    """Gold model: Calc sheet with =SUM formula + Cap Table sheet copy."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Calc"

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 28

    ws.cell(row=1, column=1, value="Cap Table FD count").font = Font(bold=True, size=13)

    ws.cell(row=3, column=1, value="Reference file")
    ws.cell(row=3, column=2, value="inputs/messy_cap_table.xlsx")

    ws.cell(row=5, column=1, value="Total FD (from reference)").font = Font(bold=True)
    ws.cell(row=5, column=2, value="=SUM('Cap Table'!C2:C10)").number_format = "#,##0"

    # Mirror the reference Cap Table as a second sheet in the model workbook.
    cap_ws = wb.create_sheet(title="Cap Table")
    _write_cap_table(cap_ws)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    recalc_xlsx(path)


def build_results_json(path: Path) -> None:
    data = {"total_fd": "Calc!B5"}
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, indent=2))


def main() -> None:
    build_reference(TASK_DIR / "inputs" / "messy_cap_table.xlsx")
    build_stub(TASK_DIR / "stubs" / "starter.xlsx")
    build_gold(TASK_DIR / "gold" / "model.xlsx")
    build_results_json(TASK_DIR / "gold" / "results.json")
    print(f"t0ref_cap_table_fd_count artifacts written. GOLD_TOTAL_FD = {GOLD_TOTAL_FD:,}")


if __name__ == "__main__":
    main()
