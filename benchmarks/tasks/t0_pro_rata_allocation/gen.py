"""Generator for t0_pro_rata_allocation.

Deterministically produces:
  stubs/starter.xlsx        — workbook with Calc sheet inputs filled, answer cells blank
  gold/model.xlsx           — same as stub but with the three pro-rata formulas filled in
  gold/results.json         — maps each output_key → "Sheet!Cell" address

Math: pool of $15M allocated pro-rata by prior dollar investment.
  allocation_i = pool * prior_i / total_prior

Given:
  pool         = 15,000,000
  nimbus_prior = 20,000,000
  arc_prior    = 10,000,000
  orbit_prior  =  5,000,000
  total_prior  = 35,000,000

Gold:
  nimbus_alloc = 15e6 * 20e6 / 35e6 = 8,571,428.5714...
  arc_alloc    = 15e6 * 10e6 / 35e6 = 4,285,714.2857...
  orbit_alloc  = 15e6 *  5e6 / 35e6 = 2,142,857.1428...
"""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

from recalc import recalc_xlsx

TASK_DIR = Path(__file__).resolve().parent

# --- parameters -------------------------------------------------------------
POOL = 15_000_000
NIMBUS_PRIOR = 20_000_000
ARC_PRIOR = 10_000_000
ORBIT_PRIOR = 5_000_000
TOTAL_PRIOR = NIMBUS_PRIOR + ARC_PRIOR + ORBIT_PRIOR

GOLD_NIMBUS = POOL * NIMBUS_PRIOR / TOTAL_PRIOR
GOLD_ARC = POOL * ARC_PRIOR / TOTAL_PRIOR
GOLD_ORBIT = POOL * ORBIT_PRIOR / TOTAL_PRIOR


# --- builders ---------------------------------------------------------------


def build_starter(path: Path, *, fill_answers: bool) -> None:
    """Build a workbook with a Calc sheet holding the pro-rata inputs.

    Layout:
      A1 Title
      A3:B3  Pool to allocate
      A4:B4  Nimbus prior inv
      A5:B5  Arc prior inv
      A6:B6  Orbit prior inv
      A7:B7  Total prior (=SUM(B4:B6))

      A10:B10 Nimbus allocation  (formula in gold, blank in stub)
      A11:B11 Arc allocation
      A12:B12 Orbit allocation
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Calc"

    hdr = Font(bold=True)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 18

    ws.cell(row=1, column=1, value="Pro-rata allocation").font = hdr

    ws.cell(row=3, column=1, value="Pool to allocate")
    ws.cell(row=3, column=2, value=POOL).number_format = "$#,##0"

    ws.cell(row=4, column=1, value="Nimbus prior inv")
    ws.cell(row=4, column=2, value=NIMBUS_PRIOR).number_format = "$#,##0"

    ws.cell(row=5, column=1, value="Arc prior inv")
    ws.cell(row=5, column=2, value=ARC_PRIOR).number_format = "$#,##0"

    ws.cell(row=6, column=1, value="Orbit prior inv")
    ws.cell(row=6, column=2, value=ORBIT_PRIOR).number_format = "$#,##0"

    ws.cell(row=7, column=1, value="Total prior")
    ws.cell(row=7, column=2, value="=SUM(B4:B6)").number_format = "$#,##0"

    ws.cell(row=9, column=1, value="Answers:").font = hdr

    ws.cell(row=10, column=1, value="Nimbus allocation")
    ws.cell(row=10, column=2).number_format = "$#,##0"
    ws.cell(row=11, column=1, value="Arc allocation")
    ws.cell(row=11, column=2).number_format = "$#,##0"
    ws.cell(row=12, column=1, value="Orbit allocation")
    ws.cell(row=12, column=2).number_format = "$#,##0"

    if fill_answers:
        ws.cell(row=10, column=2, value="=B3*B4/B7").number_format = "$#,##0"
        ws.cell(row=11, column=2, value="=B3*B5/B7").number_format = "$#,##0"
        ws.cell(row=12, column=2, value="=B3*B6/B7").number_format = "$#,##0"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def build_results_json(path: Path) -> None:
    data = {
        "nimbus_allocation": "Calc!B10",
        "arc_allocation": "Calc!B11",
        "orbit_allocation": "Calc!B12",
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, indent=2))


def main() -> None:
    build_starter(TASK_DIR / "stubs" / "starter.xlsx", fill_answers=False)
    gold = TASK_DIR / "gold" / "model.xlsx"
    build_starter(gold, fill_answers=True)
    recalc_xlsx(gold)
    build_results_json(TASK_DIR / "gold" / "results.json")
    print(f"GOLD_NIMBUS = {GOLD_NIMBUS:,.4f}")
    print(f"GOLD_ARC    = {GOLD_ARC:,.4f}")
    print(f"GOLD_ORBIT  = {GOLD_ORBIT:,.4f}")


if __name__ == "__main__":
    main()
