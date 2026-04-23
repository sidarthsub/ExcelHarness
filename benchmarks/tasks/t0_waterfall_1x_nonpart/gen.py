"""Generator for t0_waterfall_1x_nonpart.

Deterministically produces:
  stubs/starter.xlsx       — workbook with inputs filled, answer cells blank
  gold/model.xlsx          — same as stub but with correct formulas
  gold/results.json        — semantic keys → Sheet!Cell addresses

Math: 1x non-participating preferred liquidation waterfall.

  pref_amount    = investment × multiple
  convert_amount = ownership_pct × exit_value
  lp_proceeds    = MAX(pref_amount, convert_amount)
  common_proceeds = exit_value − lp_proceeds

With the given parameters:
  pref_amount    = $12M × 1.0       = $12,000,000
  convert_amount = 0.30 × $45M      = $13,500,000
  lp_proceeds    = MAX(12M, 13.5M)  = $13,500,000   (conversion beats pref)
  common_proceeds = $45M − $13.5M   = $31,500,000
"""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

from benchmarks.recalc import recalc_xlsx

TASK_DIR = Path(__file__).resolve().parent

# --- parameters -------------------------------------------------------------
LP_INVESTMENT = 12_000_000
LP_OWNERSHIP = 0.30
PREF_MULTIPLE = 1.0
EXIT_VALUE = 45_000_000

GOLD_PREF = LP_INVESTMENT * PREF_MULTIPLE          # 12_000_000
GOLD_CONVERT = LP_OWNERSHIP * EXIT_VALUE            # 13_500_000
GOLD_LP = max(GOLD_PREF, GOLD_CONVERT)              # 13_500_000
GOLD_COMMON = EXIT_VALUE - GOLD_LP                  # 31_500_000


# --- builders ---------------------------------------------------------------


def build_starter(path: Path, *, fill_formulas: bool) -> None:
    """Build the Calc sheet. If fill_formulas, populate the answer cells
    (B10:B13) with formulas referencing B3:B6. Otherwise leave blank.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Calc"

    hdr = Font(bold=True)
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 18

    ws.cell(row=1, column=1, value="1x non-participating waterfall").font = hdr

    ws.cell(row=3, column=1, value="LP investment")
    ws.cell(row=3, column=2, value=LP_INVESTMENT).number_format = "#,##0"

    ws.cell(row=4, column=1, value="LP ownership (%)")
    ws.cell(row=4, column=2, value=LP_OWNERSHIP).number_format = "0.00%"

    ws.cell(row=5, column=1, value="Pref multiple")
    ws.cell(row=5, column=2, value=PREF_MULTIPLE).number_format = "0.00"

    ws.cell(row=6, column=1, value="Exit value")
    ws.cell(row=6, column=2, value=EXIT_VALUE).number_format = "#,##0"

    ws.cell(row=9, column=1, value="Answers:").font = hdr

    ws.cell(row=10, column=1, value="Preference amount")
    ws.cell(row=11, column=1, value="Conversion amount")
    ws.cell(row=12, column=1, value="LP actual proceeds")
    ws.cell(row=13, column=1, value="Common proceeds")

    if fill_formulas:
        ws.cell(row=10, column=2, value="=B3*B5")
        ws.cell(row=11, column=2, value="=B4*B6")
        ws.cell(row=12, column=2, value="=MAX(B10,B11)")
        ws.cell(row=13, column=2, value="=B6-B12")

    for r in (10, 11, 12, 13):
        ws.cell(row=r, column=2).number_format = "#,##0"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def build_results_json(path: Path) -> None:
    data = {
        "pref_amount": "Calc!B10",
        "convert_amount": "Calc!B11",
        "lp_proceeds": "Calc!B12",
        "common_proceeds": "Calc!B13",
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, indent=2))


def main() -> None:
    build_starter(TASK_DIR / "stubs" / "starter.xlsx", fill_formulas=False)
    gold = TASK_DIR / "gold" / "model.xlsx"
    build_starter(gold, fill_formulas=True)
    recalc_xlsx(gold)
    build_results_json(TASK_DIR / "gold" / "results.json")
    print("t0_waterfall_1x_nonpart artifacts written.")
    print(f"  pref_amount     = {GOLD_PREF:,.2f}")
    print(f"  convert_amount  = {GOLD_CONVERT:,.2f}")
    print(f"  lp_proceeds     = {GOLD_LP:,.2f}")
    print(f"  common_proceeds = {GOLD_COMMON:,.2f}")


if __name__ == "__main__":
    main()
