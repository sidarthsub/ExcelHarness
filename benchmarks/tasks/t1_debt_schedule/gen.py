"""Generator for t1_debt_schedule.

Task: build a 60-month debt amortization schedule on a single sheet
named `Amortization Schedule`. Principal $1,000,000, annual rate 6.0%,
60 monthly periods. Payment computed via Excel's PMT.

The stub is a minimal workbook with a ReadMe sheet. The gold is the
canonical amortization schedule with parameters + 60 monthly rows +
summary totals, all driven by formulas that reference the parameter
block.

Grading is per-cell numeric against expected values + formula-driven
for derived fields (payment, interest, principal, ending balance, totals).
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from recalc import recalc_xlsx

TASK_DIR = Path(__file__).resolve().parent

# --- parameters -------------------------------------------------------------
PRINCIPAL = 1_000_000
ANNUAL_RATE = 0.06
N_PERIODS = 60

# Table layout constants
PARAM_ROWS = {
    "principal": 2,
    "rate": 3,
    "term": 4,
    "payment": 5,
}
HEADER_ROW = 7
FIRST_DATA_ROW = 8
LAST_DATA_ROW = FIRST_DATA_ROW + N_PERIODS - 1  # 67
SUMMARY_START_ROW = LAST_DATA_ROW + 3  # 70


def _header_fill(cell) -> None:
    cell.font = Font(bold=True)
    cell.fill = PatternFill("solid", fgColor="EEEEEE")


def build_schedule_sheet(ws, *, fill_formulas: bool) -> None:
    ws.title = "Amortization Schedule"
    ws.column_dimensions["A"].width = 8
    for col in ("B", "C", "D", "E", "F"):
        ws.column_dimensions[col].width = 14

    # Row 1: title
    title = ws.cell(row=1, column=1, value="Amortization Schedule")
    title.font = Font(bold=True, size=13)

    # Rows 2-5: parameter block
    ws.cell(row=2, column=1, value="Principal")
    ws.cell(row=2, column=2, value=PRINCIPAL).number_format = "$#,##0"

    ws.cell(row=3, column=1, value="Annual rate")
    ws.cell(row=3, column=2, value=ANNUAL_RATE).number_format = "0.00%"

    ws.cell(row=4, column=1, value="Term (months)")
    ws.cell(row=4, column=2, value=N_PERIODS).number_format = "0"

    ws.cell(row=5, column=1, value="Monthly payment")
    if fill_formulas:
        ws.cell(row=5, column=2, value="=PMT(B3/12,B4,-B2)").number_format = "$#,##0.00"
    else:
        ws.cell(row=5, column=2).number_format = "$#,##0.00"

    # Row 7: headers
    headers = [
        "Month",
        "Beginning Balance",
        "Payment",
        "Interest",
        "Principal",
        "Ending Balance",
    ]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=HEADER_ROW, column=i, value=h)
        _header_fill(c)
        if i == 1:
            c.alignment = Alignment(horizontal="center")

    # Rows 8..67: monthly schedule
    for t in range(N_PERIODS):
        r = FIRST_DATA_ROW + t
        # Month
        month_cell = ws.cell(row=r, column=1)
        if fill_formulas:
            if t == 0:
                month_cell.value = 1
            else:
                month_cell.value = f"=A{r-1}+1"
        month_cell.number_format = "0"
        month_cell.alignment = Alignment(horizontal="center")

        # Beginning balance
        beg = ws.cell(row=r, column=2)
        if fill_formulas:
            if t == 0:
                beg.value = "=$B$2"
            else:
                beg.value = f"=F{r-1}"
        beg.number_format = "$#,##0.00"

        # Payment
        pay = ws.cell(row=r, column=3)
        if fill_formulas:
            pay.value = "=$B$5"
        pay.number_format = "$#,##0.00"

        # Interest
        interest = ws.cell(row=r, column=4)
        if fill_formulas:
            interest.value = f"=B{r}*$B$3/12"
        interest.number_format = "$#,##0.00"

        # Principal
        principal = ws.cell(row=r, column=5)
        if fill_formulas:
            principal.value = f"=C{r}-D{r}"
        principal.number_format = "$#,##0.00"

        # Ending balance
        ending = ws.cell(row=r, column=6)
        if fill_formulas:
            ending.value = f"=B{r}-E{r}"
        ending.number_format = "$#,##0.00"

    # Summary section (rows 70-72)
    r = SUMMARY_START_ROW
    lbl1 = ws.cell(row=r, column=1, value="Total interest paid")
    lbl1.font = Font(bold=True)
    if fill_formulas:
        ws.cell(row=r, column=2, value=f"=SUM(D{FIRST_DATA_ROW}:D{LAST_DATA_ROW})").number_format = "$#,##0.00"
    else:
        ws.cell(row=r, column=2).number_format = "$#,##0.00"

    r += 1
    lbl2 = ws.cell(row=r, column=1, value="Total principal paid")
    lbl2.font = Font(bold=True)
    if fill_formulas:
        ws.cell(row=r, column=2, value=f"=SUM(E{FIRST_DATA_ROW}:E{LAST_DATA_ROW})").number_format = "$#,##0.00"
    else:
        ws.cell(row=r, column=2).number_format = "$#,##0.00"

    r += 1
    lbl3 = ws.cell(row=r, column=1, value="Final balance")
    lbl3.font = Font(bold=True)
    if fill_formulas:
        ws.cell(row=r, column=2, value=f"=F{LAST_DATA_ROW}").number_format = "$#,##0.00"
    else:
        ws.cell(row=r, column=2).number_format = "$#,##0.00"


def build_stub(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ReadMe"
    ws.cell(
        row=1,
        column=1,
        value="Build Amortization Schedule per brief and inputs/schedule_format.md.",
    )
    ws.column_dimensions["A"].width = 80
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def build_gold(path: Path) -> None:
    wb = openpyxl.Workbook()
    build_schedule_sheet(wb.active, fill_formulas=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    recalc_xlsx(path)


def build_results_json(path: Path) -> None:
    """Gold results.json: each semantic key mapped to its Sheet!Cell address.

    Layout built by build_schedule_sheet(fill_formulas=True):
      B5   Monthly payment (PMT)
      D8   Month 1 interest
      E8   Month 1 principal
      F67  Month 60 ending balance
      B70  Total interest paid
      B71  Total principal paid
    """
    import json

    sheet = "'Amortization Schedule'"
    data = {
        "monthly_payment": f"{sheet}!B5",
        "month_1_interest": f"{sheet}!D8",
        "month_1_principal": f"{sheet}!E8",
        "month_60_ending_balance": f"{sheet}!F67",
        "total_interest_paid": f"{sheet}!B70",
        "total_principal_paid": f"{sheet}!B71",
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, indent=2))


def main() -> None:
    build_stub(TASK_DIR / "stubs" / "starter.xlsx")
    build_gold(TASK_DIR / "gold" / "model.xlsx")
    build_results_json(TASK_DIR / "gold" / "results.json")
    print("t1_debt_schedule artifacts written.")


if __name__ == "__main__":
    main()
