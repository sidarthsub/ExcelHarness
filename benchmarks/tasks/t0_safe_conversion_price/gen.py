"""Generator for t0_safe_conversion_price.

Deterministically produces:
  stubs/starter.xlsx      — workbook with Calc sheet inputs filled, answer cells blank
  gold/model.xlsx         — same as stub but with correct formulas in B10:B14
  gold/results.json       — maps output_keys to their Sheet!Cell addresses

Math: SAFE conversion at a priced round.

  headline_price   = round_pre_money / pre_round_FD
  cap_price        = valuation_cap   / pre_round_FD
  discount_price   = (1 - discount) * headline_price
  conversion_price = MIN(cap_price, discount_price)
  shares_issued    = SAFE_investment / conversion_price

With the defaults below:
  headline_price   = 20M / 5M      = 4.00
  cap_price        = 15M / 5M      = 3.00
  discount_price   = 0.80 * 4.00   = 3.20
  conversion_price = MIN(3.00,3.20) = 3.00
  shares_issued    = 3M / 3.00     = 1,000,000
"""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

from benchmarks.recalc import recalc_xlsx

TASK_DIR = Path(__file__).resolve().parent

# --- parameters -------------------------------------------------------------
SAFE_INVESTMENT = 3_000_000
VALUATION_CAP = 15_000_000
DISCOUNT = 0.20
PRE_ROUND_FD = 5_000_000
ROUND_PRE_MONEY = 20_000_000

# Expected values (for sanity / reference)
GOLD_HEADLINE = ROUND_PRE_MONEY / PRE_ROUND_FD                    # 4.0
GOLD_CAP_PRICE = VALUATION_CAP / PRE_ROUND_FD                     # 3.0
GOLD_DISCOUNT_PRICE = (1 - DISCOUNT) * GOLD_HEADLINE               # 3.2
GOLD_CONVERSION = min(GOLD_CAP_PRICE, GOLD_DISCOUNT_PRICE)         # 3.0
GOLD_SHARES = SAFE_INVESTMENT / GOLD_CONVERSION                    # 1_000_000


# --- builders ---------------------------------------------------------------


def build_calc_sheet(ws, *, fill_formulas: bool) -> None:
    ws.title = "Calc"
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 18

    title = ws.cell(row=1, column=1, value="SAFE Conversion Price")
    title.font = Font(bold=True, size=13)

    # Inputs
    ws.cell(row=3, column=1, value="SAFE investment")
    ws.cell(row=3, column=2, value=SAFE_INVESTMENT).number_format = "$#,##0"

    ws.cell(row=4, column=1, value="Valuation cap")
    ws.cell(row=4, column=2, value=VALUATION_CAP).number_format = "$#,##0"

    ws.cell(row=5, column=1, value="Discount (decimal)")
    ws.cell(row=5, column=2, value=DISCOUNT).number_format = "0.00%"

    ws.cell(row=6, column=1, value="Pre-round FD")
    ws.cell(row=6, column=2, value=PRE_ROUND_FD).number_format = "#,##0"

    ws.cell(row=7, column=1, value="Round pre-money")
    ws.cell(row=7, column=2, value=ROUND_PRE_MONEY).number_format = "$#,##0"

    # Answer cells
    ws.cell(row=10, column=1, value="Headline price")
    ws.cell(row=11, column=1, value="Cap price")
    ws.cell(row=12, column=1, value="Discount price")
    ws.cell(row=13, column=1, value="Conversion price")
    ws.cell(row=14, column=1, value="Shares issued")

    if fill_formulas:
        ws.cell(row=10, column=2, value="=B7/B6").number_format = "$#,##0.0000"
        ws.cell(row=11, column=2, value="=B4/B6").number_format = "$#,##0.0000"
        ws.cell(row=12, column=2, value="=(1-B5)*B10").number_format = "$#,##0.0000"
        ws.cell(row=13, column=2, value="=MIN(B11,B12)").number_format = "$#,##0.0000"
        ws.cell(row=14, column=2, value="=B3/B13").number_format = "#,##0"
    else:
        ws.cell(row=10, column=2).number_format = "$#,##0.0000"
        ws.cell(row=11, column=2).number_format = "$#,##0.0000"
        ws.cell(row=12, column=2).number_format = "$#,##0.0000"
        ws.cell(row=13, column=2).number_format = "$#,##0.0000"
        ws.cell(row=14, column=2).number_format = "#,##0"


def build_stub(path: Path) -> None:
    wb = openpyxl.Workbook()
    build_calc_sheet(wb.active, fill_formulas=False)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def build_gold(path: Path) -> None:
    wb = openpyxl.Workbook()
    build_calc_sheet(wb.active, fill_formulas=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    recalc_xlsx(path)


def build_results_json(path: Path) -> None:
    data = {
        "headline_price": "Calc!B10",
        "cap_price": "Calc!B11",
        "discount_price": "Calc!B12",
        "conversion_price": "Calc!B13",
        "shares_issued": "Calc!B14",
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, indent=2))


def main() -> None:
    build_stub(TASK_DIR / "stubs" / "starter.xlsx")
    build_gold(TASK_DIR / "gold" / "model.xlsx")
    build_results_json(TASK_DIR / "gold" / "results.json")
    print(
        "t0_safe_conversion_price artifacts written. "
        f"headline={GOLD_HEADLINE}, cap={GOLD_CAP_PRICE}, "
        f"discount={GOLD_DISCOUNT_PRICE}, conversion={GOLD_CONVERSION}, "
        f"shares={GOLD_SHARES}"
    )


if __name__ == "__main__":
    main()
