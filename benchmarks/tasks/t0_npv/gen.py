"""Generator for t0_npv.

Deterministically produces:
  stubs/starter.xlsx  - workbook with Calc sheet inputs filled, answer blank
  gold/model.xlsx     - same as stub but with the correct NPV formula
  gold/results.json   - maps semantic key `npv` to Calc!B13
  gold/grading.yaml   - rubric (hardcoded, with the computed expected value)

Math: NPV of a cash flow series.
  Year 0 investment:  -10,000,000
  Years 1-5 inflows:   3,000,000 each
  Discount rate:       10%

  NPV = -10M + sum_{t=1..5} 3M / (1.10)^t
      = -10M + 3M * ((1 - 1.10^-5) / 0.10)
      ~= 1,372,360.31

Excel convention: `NPV(rate, values)` treats `values` as starting at t=1, so
the canonical gold formula is:
  =B10 + NPV(B3, B5:B9)
"""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

from benchmarks.recalc import recalc_xlsx

TASK_DIR = Path(__file__).resolve().parent

# --- parameters -------------------------------------------------------------
DISCOUNT_RATE = 0.10
Y0_CASH_FLOW = -10_000_000
YEARLY_INFLOW = 3_000_000
N_YEARS = 5

# Expected NPV, computed at gen time so grading.yaml stays in sync.
GOLD_NPV = sum(
    cf / (1 + DISCOUNT_RATE) ** i
    for i, cf in enumerate([Y0_CASH_FLOW] + [YEARLY_INFLOW] * N_YEARS)
)
# ~= 1,372,360.3082...


# --- builders ---------------------------------------------------------------


def build_starter(path: Path, *, fill_answer: bool) -> None:
    """Build a workbook with a single Calc sheet.

    Layout:
      A1  "NPV calculation"
      A3  "Discount rate"    B3  0.10
      A5  "Year 1"           B5  3,000,000
      A6  "Year 2"           B6  3,000,000
      A7  "Year 3"           B7  3,000,000
      A8  "Year 4"           B8  3,000,000
      A9  "Year 5"           B9  3,000,000
      A10 "Year 0"           B10 -10,000,000
      A13 "NPV"              B13 <formula if fill_answer>
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Calc"

    hdr = Font(bold=True)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 18

    ws.cell(row=1, column=1, value="NPV calculation").font = hdr

    ws.cell(row=3, column=1, value="Discount rate")
    ws.cell(row=3, column=2, value=DISCOUNT_RATE).number_format = "0.00%"

    for i in range(N_YEARS):
        year_row = 5 + i
        ws.cell(row=year_row, column=1, value=f"Year {i + 1}")
        ws.cell(row=year_row, column=2, value=YEARLY_INFLOW).number_format = "$#,##0"

    ws.cell(row=10, column=1, value="Year 0")
    ws.cell(row=10, column=2, value=Y0_CASH_FLOW).number_format = "$#,##0"

    ws.cell(row=13, column=1, value="NPV").font = hdr
    if fill_answer:
        ws.cell(row=13, column=2, value="=B10+NPV(B3,B5:B9)")
    ws.cell(row=13, column=2).number_format = "$#,##0.00"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def build_results_json(path: Path) -> None:
    data = {"npv": "Calc!B13"}
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, indent=2))


def build_grading_yaml(path: Path) -> None:
    """Emit grading.yaml with the computed expected NPV baked in."""
    content = f"""recalc: true
checks:
  - id: calc_sheet_exists
    type: sheet_exists
    sheet: Calc
    weight: 0.3

  - id: npv_value
    type: output_key
    key: npv
    expected: {GOLD_NPV!r}
    tolerance_rel: 1.0e-4
    require_formula: true
    weight: 1.5
    description: "NPV of cash flows must match and be formula-driven."
"""
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content)


def main() -> None:
    build_starter(TASK_DIR / "stubs" / "starter.xlsx", fill_answer=False)

    gold = TASK_DIR / "gold" / "model.xlsx"
    build_starter(gold, fill_answer=True)
    recalc_xlsx(gold)

    build_results_json(TASK_DIR / "gold" / "results.json")
    build_grading_yaml(TASK_DIR / "gold" / "grading.yaml")

    print(f"GOLD_NPV = {GOLD_NPV:,.4f}")
    print("t0_npv artifacts written.")


if __name__ == "__main__":
    main()
