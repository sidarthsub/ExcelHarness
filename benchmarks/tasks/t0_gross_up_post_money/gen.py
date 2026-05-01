"""Generator for t0_gross_up_post_money.

Deterministically produces:
  stubs/starter.xlsx  - workbook with Calc sheet inputs filled, answers blank
  gold/model.xlsx     - same as stub but with the correct formulas
  gold/results.json   - maps semantic keys to Calc!B9..B13
  gold/grading.yaml   - rubric (hardcoded, with the computed expected values)

Math: solve for pre-money given a target post-money %.
  target_post_fd  = holder_shares / target_pct
  new_shares      = target_post_fd - pre_fd
  share_price     = round_size / new_shares
  pre_money       = pre_fd * share_price
  holder_pct_check = holder_shares / target_post_fd

Parameters:
  holder_shares   = 3,000,000
  pre_fd          = 5,000,000
  target_post_pct = 0.20
  round_size      = 15,000,000

Expected:
  target_post_fd   = 15,000,000
  new_shares       = 10,000,000
  share_price      = 1.50
  pre_money        = 7,500,000
  holder_pct_check = 0.20
"""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

from recalc import recalc_xlsx

TASK_DIR = Path(__file__).resolve().parent

# --- parameters -------------------------------------------------------------
HOLDER_SHARES = 3_000_000
PRE_FD = 5_000_000
TARGET_POST_PCT = 0.20
ROUND_SIZE = 15_000_000

# Expected values, computed at gen time so grading.yaml stays in sync.
GOLD_TARGET_POST_FD = HOLDER_SHARES / TARGET_POST_PCT
GOLD_NEW_SHARES = GOLD_TARGET_POST_FD - PRE_FD
GOLD_SHARE_PRICE = ROUND_SIZE / GOLD_NEW_SHARES
GOLD_PRE_MONEY = PRE_FD * GOLD_SHARE_PRICE
GOLD_HOLDER_PCT_CHECK = HOLDER_SHARES / GOLD_TARGET_POST_FD


# --- builders ---------------------------------------------------------------


def build_starter(path: Path, *, fill_answer: bool) -> None:
    """Build a workbook with a single Calc sheet.

    Layout:
      A1  "Gross-up to target post-money %"
      A3  "Holder shares"     B3  3,000,000
      A4  "Pre-round FD"      B4  5,000,000
      A5  "Target holder %"   B5  0.20
      A6  "Round size"        B6  15,000,000
      A9  "Target post-FD"    B9  =B3/B5
      A10 "New shares"        B10 =B9-B4
      A11 "Share price"       B11 =B6/B10
      A12 "Pre-money"         B12 =B4*B11
      A13 "Check holder %"    B13 =B3/B9
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Calc"

    hdr = Font(bold=True)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 18

    ws.cell(row=1, column=1, value="Gross-up to target post-money %").font = hdr

    ws.cell(row=3, column=1, value="Holder shares")
    ws.cell(row=3, column=2, value=HOLDER_SHARES).number_format = "#,##0"

    ws.cell(row=4, column=1, value="Pre-round FD")
    ws.cell(row=4, column=2, value=PRE_FD).number_format = "#,##0"

    ws.cell(row=5, column=1, value="Target holder %")
    ws.cell(row=5, column=2, value=TARGET_POST_PCT).number_format = "0.00%"

    ws.cell(row=6, column=1, value="Round size")
    ws.cell(row=6, column=2, value=ROUND_SIZE).number_format = "$#,##0"

    ws.cell(row=9, column=1, value="Target post-FD")
    ws.cell(row=10, column=1, value="New shares")
    ws.cell(row=11, column=1, value="Share price")
    ws.cell(row=12, column=1, value="Pre-money")
    ws.cell(row=13, column=1, value="Check holder %")

    ws.cell(row=9, column=2).number_format = "#,##0"
    ws.cell(row=10, column=2).number_format = "#,##0"
    ws.cell(row=11, column=2).number_format = "$#,##0.00"
    ws.cell(row=12, column=2).number_format = "$#,##0"
    ws.cell(row=13, column=2).number_format = "0.00%"

    if fill_answer:
        ws.cell(row=9, column=2, value="=B3/B5")
        ws.cell(row=10, column=2, value="=B9-B4")
        ws.cell(row=11, column=2, value="=B6/B10")
        ws.cell(row=12, column=2, value="=B4*B11")
        ws.cell(row=13, column=2, value="=B3/B9")

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def build_results_json(path: Path) -> None:
    data = {
        "target_post_fd": "Calc!B9",
        "new_shares_issued": "Calc!B10",
        "share_price": "Calc!B11",
        "pre_money": "Calc!B12",
        "holder_pct_check": "Calc!B13",
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, indent=2))


def build_grading_yaml(path: Path) -> None:
    """Emit grading.yaml with the computed expected values baked in."""
    content = f"""recalc: true
checks:
  - id: calc_sheet_exists
    type: sheet_exists
    sheet: Calc
    weight: 0.3

  - id: target_post_fd_value
    type: output_key
    key: target_post_fd
    expected: {GOLD_TARGET_POST_FD!r}
    tolerance_rel: 1.0e-6
    require_formula: true
    weight: 0.7
    description: "Target post-round fully diluted share count."

  - id: new_shares_issued_value
    type: output_key
    key: new_shares_issued
    expected: {GOLD_NEW_SHARES!r}
    tolerance_rel: 1.0e-6
    require_formula: true
    weight: 0.7
    description: "New investor shares issued in the round."

  - id: share_price_value
    type: output_key
    key: share_price
    expected: {GOLD_SHARE_PRICE!r}
    tolerance_rel: 1.0e-6
    require_formula: true
    weight: 0.8
    description: "Per-share price for the round."

  - id: pre_money_value
    type: output_key
    key: pre_money
    expected: {GOLD_PRE_MONEY!r}
    tolerance_rel: 1.0e-6
    require_formula: true
    weight: 1.2
    description: "Pre-money valuation implied by the target holder %."

  - id: holder_pct_check_value
    type: output_key
    key: holder_pct_check
    expected: {GOLD_HOLDER_PCT_CHECK!r}
    tolerance_rel: 1.0e-6
    require_formula: true
    weight: 0.5
    description: "Sanity check: holder % post-round equals target."
"""
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content)


def main() -> None:
    build_starter(TASK_DIR / "stubs" / "starter.xlsx", fill_answer=False)

    gold = TASK_DIR / "gold" / "model.xlsx"
    build_starter(gold, fill_answer=True)
    recalc_xlsx(gold)

    build_results_json(TASK_DIR / "gold" / "results.json")
    build_grading_yaml(TASK_DIR / "gold" / "grading.yaml")

    print(f"GOLD_TARGET_POST_FD   = {GOLD_TARGET_POST_FD:,.4f}")
    print(f"GOLD_NEW_SHARES       = {GOLD_NEW_SHARES:,.4f}")
    print(f"GOLD_SHARE_PRICE      = {GOLD_SHARE_PRICE:,.4f}")
    print(f"GOLD_PRE_MONEY        = {GOLD_PRE_MONEY:,.4f}")
    print(f"GOLD_HOLDER_PCT_CHECK = {GOLD_HOLDER_PCT_CHECK:,.6f}")
    print("t0_gross_up_post_money artifacts written.")


if __name__ == "__main__":
    main()
