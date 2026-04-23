"""Generator for t0ref_historical_cagr.

Deterministically produces:
  inputs/historical_revenue.xlsx — reference workbook with a Revenue sheet
  stubs/starter.xlsx             — Calc sheet laid out, answer cells blank
  gold/model.xlsx                — Revenue imported + Calc filled with formulas
  gold/results.json              — semantic key -> Sheet!Cell address map

Math:
  3-year CAGR ending 2025 = (end/start)^(1/n) - 1
    start = 2022 revenue (55M), end = 2025 revenue (130M), n = 3
  4-year CAGR (2021 -> 2025) = (130M / 40M)^(1/4) - 1
  Total growth multiple     = 130M / 40M = 3.25x
"""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from benchmarks.recalc import recalc_xlsx

TASK_DIR = Path(__file__).resolve().parent

# --- parameters -------------------------------------------------------------
# Revenue rows: (year, revenue). First row is header; data starts at row 2.
REVENUE_ROWS = [
    (2021, 40_000_000),
    (2022, 55_000_000),
    (2023, 75_000_000),
    (2024, 100_000_000),
    (2025, 130_000_000),
]

# Index helpers (1-based rows). Header on row 1; first data row is row 2.
YEAR_TO_ROW = {year: idx + 2 for idx, (year, _) in enumerate(REVENUE_ROWS)}
# -> {2021: 2, 2022: 3, 2023: 4, 2024: 5, 2025: 6}

START_YEAR = 2022
END_YEAR = 2025
BASELINE_YEAR = 2021  # for the 4-year CAGR and growth multiple


# --- gold math --------------------------------------------------------------


def compute_gold() -> dict:
    rev = dict(REVENUE_ROWS)
    start_rev = rev[START_YEAR]
    end_rev = rev[END_YEAR]
    baseline_rev = rev[BASELINE_YEAR]
    years = END_YEAR - START_YEAR  # 3
    cagr_3yr = (end_rev / start_rev) ** (1 / years) - 1
    cagr_4yr = (end_rev / baseline_rev) ** (1 / (END_YEAR - BASELINE_YEAR)) - 1
    growth_multiple = end_rev / baseline_rev
    return {
        "start_rev": start_rev,
        "end_rev": end_rev,
        "years": years,
        "cagr_3yr": cagr_3yr,
        "cagr_4yr": cagr_4yr,
        "growth_multiple": growth_multiple,
    }


GOLD = compute_gold()


# --- builders ---------------------------------------------------------------


def _build_revenue_sheet(ws) -> None:
    ws.title = "Revenue"
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 16

    hdr_font = Font(bold=True)
    hdr_fill = PatternFill("solid", fgColor="E8E8E8")
    for col, h in enumerate(["Year", "Revenue"], 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = hdr_font
        c.fill = hdr_fill
        c.alignment = Alignment(horizontal="center")

    for idx, (year, revenue) in enumerate(REVENUE_ROWS, 2):
        ws.cell(row=idx, column=1, value=year)
        ws.cell(row=idx, column=2, value=revenue).number_format = "#,##0"


def build_reference(path: Path) -> None:
    wb = openpyxl.Workbook()
    _build_revenue_sheet(wb.active)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    recalc_xlsx(path)


def _build_calc_sheet(ws, *, fill_formulas: bool) -> None:
    ws.title = "Calc"
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 20

    hdr = Font(bold=True)
    ws.cell(row=1, column=1, value="Historical revenue CAGR").font = Font(bold=True, size=13)

    ws.cell(row=3, column=1, value="Reference file")
    ws.cell(row=3, column=2, value="inputs/historical_revenue.xlsx")

    ws.cell(row=5, column=1, value="Start revenue (2022)")
    ws.cell(row=6, column=1, value="End revenue (2025)")
    ws.cell(row=7, column=1, value="Years elapsed")

    ws.cell(row=9, column=1, value="Answers:").font = hdr
    ws.cell(row=10, column=1, value="3-year CAGR (ending 2025)")
    ws.cell(row=11, column=1, value="4-year CAGR (2021 -> 2025)")
    ws.cell(row=12, column=1, value="Total growth multiple")

    start_row = YEAR_TO_ROW[START_YEAR]        # 3
    end_row = YEAR_TO_ROW[END_YEAR]            # 6
    baseline_row = YEAR_TO_ROW[BASELINE_YEAR]  # 2

    if fill_formulas:
        ws.cell(row=5, column=2, value=f"=Revenue!B{start_row}").number_format = "#,##0"
        ws.cell(row=6, column=2, value=f"=Revenue!B{end_row}").number_format = "#,##0"
        ws.cell(row=7, column=2, value=END_YEAR - START_YEAR)
        ws.cell(row=10, column=2, value="=(B6/B5)^(1/B7)-1").number_format = "0.00%"
        ws.cell(row=11, column=2, value=f"=(B6/Revenue!B{baseline_row})^(1/4)-1").number_format = "0.00%"
        ws.cell(row=12, column=2, value=f"=B6/Revenue!B{baseline_row}").number_format = "0.00"
    else:
        ws.cell(row=5, column=2).number_format = "#,##0"
        ws.cell(row=6, column=2).number_format = "#,##0"
        ws.cell(row=10, column=2).number_format = "0.00%"
        ws.cell(row=11, column=2).number_format = "0.00%"
        ws.cell(row=12, column=2).number_format = "0.00"


def build_stub(path: Path) -> None:
    wb = openpyxl.Workbook()
    _build_calc_sheet(wb.active, fill_formulas=False)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def build_gold(path: Path) -> None:
    wb = openpyxl.Workbook()
    # Sheet 1: Calc (default active)
    _build_calc_sheet(wb.active, fill_formulas=True)
    # Sheet 2: Revenue (copied from the reference layout)
    ws_rev = wb.create_sheet(title="Revenue")
    _build_revenue_sheet(ws_rev)
    # Keep Calc first so the layout is obvious
    wb.move_sheet("Calc", offset=0)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    recalc_xlsx(path)


def build_results_json(path: Path) -> None:
    data = {
        "start_revenue": "Calc!B5",
        "end_revenue": "Calc!B6",
        "cagr_3yr": "Calc!B10",
        "cagr_4yr": "Calc!B11",
        "growth_multiple": "Calc!B12",
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, indent=2))


def main() -> None:
    build_reference(TASK_DIR / "inputs" / "historical_revenue.xlsx")
    build_stub(TASK_DIR / "stubs" / "starter.xlsx")
    build_gold(TASK_DIR / "gold" / "model.xlsx")
    build_results_json(TASK_DIR / "gold" / "results.json")
    print("t0ref_historical_cagr artifacts written.")
    print(f"  start_rev       = {GOLD['start_rev']:,}")
    print(f"  end_rev         = {GOLD['end_rev']:,}")
    print(f"  years           = {GOLD['years']}")
    print(f"  cagr_3yr        = {GOLD['cagr_3yr']:.10f}  ({GOLD['cagr_3yr']*100:.2f}%)")
    print(f"  cagr_4yr        = {GOLD['cagr_4yr']:.10f}  ({GOLD['cagr_4yr']*100:.2f}%)")
    print(f"  growth_multiple = {GOLD['growth_multiple']}")


if __name__ == "__main__":
    main()
