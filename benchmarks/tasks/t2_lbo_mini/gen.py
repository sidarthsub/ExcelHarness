"""Generator for t2_lbo_mini.

Three-sheet mini LBO:
  1. Sources and Uses — standard S&U with sponsor equity as plug
  2. Operating Model  — 5-year projection, TLB sweeps excess FCF
  3. Returns          — exit EV, net debt, MOIC, IRR

Math (iterative year-over-year for TLB balance).
"""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from recalc import recalc_xlsx

TASK_DIR = Path(__file__).resolve().parent

# --- parameters -------------------------------------------------------------
EV = 500_000_000
FEES = 10_000_000
TLB0 = 250_000_000
TLB_RATE = 0.07
MEZZ = 50_000_000
MEZZ_RATE = 0.10
MEZZ_INTEREST = 5_000_000  # flat

REV0 = 250_000_000
REV_GROWTH = 0.06
EBITDA0 = 50_000_000
EBITDA_GROWTH = 0.08
DA_FLAT = 12_000_000
TAX = 0.25
CAPEX_PCT = 0.04
MAND_AMORT = 2_500_000  # 1% × 250M

EXIT_MULT = 11.0
HOLD = 5


# --- gold math --------------------------------------------------------------


def compute_gold() -> dict:
    total_uses = EV + FEES
    sponsor_eq = total_uses - TLB0 - MEZZ
    years = []
    tlb_beg = TLB0
    for t in range(1, HOLD + 1):
        rev = REV0 * ((1 + REV_GROWTH) ** t)
        ebitda = EBITDA0 * ((1 + EBITDA_GROWTH) ** t)
        ebit = ebitda - DA_FLAT
        int_tlb = TLB_RATE * tlb_beg
        pretax = ebit - int_tlb - MEZZ_INTEREST
        tax = max(0.0, pretax) * TAX
        ni = pretax - tax
        capex = CAPEX_PCT * rev
        fcf = ni + DA_FLAT - capex
        sweep = max(0.0, min(fcf - MAND_AMORT, tlb_beg - MAND_AMORT))
        tlb_end = tlb_beg - MAND_AMORT - sweep
        years.append({
            "t": t, "rev": rev, "ebitda": ebitda, "ebit": ebit,
            "tlb_beg": tlb_beg, "int_tlb": int_tlb, "int_mezz": MEZZ_INTEREST,
            "pretax": pretax, "tax": tax, "ni": ni, "capex": capex,
            "fcf": fcf, "mand": MAND_AMORT, "sweep": sweep, "tlb_end": tlb_end,
            "mezz_end": MEZZ,
        })
        tlb_beg = tlb_end
    ebitda_y5 = years[-1]["ebitda"]
    tlb_y5_end = years[-1]["tlb_end"]
    exit_ev = EXIT_MULT * ebitda_y5
    net_debt = tlb_y5_end + MEZZ
    exit_equity = exit_ev - net_debt
    moic = exit_equity / sponsor_eq
    irr = moic ** (1 / HOLD) - 1
    return {
        "total_uses": total_uses,
        "total_sources": total_uses,
        "sponsor_eq": sponsor_eq,
        "years": years,
        "ebitda_y5": ebitda_y5,
        "y1_fcf": years[0]["fcf"],
        "y5_fcf": years[-1]["fcf"],
        "tlb_y5_end": tlb_y5_end,
        "exit_ev": exit_ev,
        "net_debt": net_debt,
        "exit_equity": exit_equity,
        "moic": moic,
        "irr": irr,
    }


GOLD = compute_gold()


# --- builders ---------------------------------------------------------------


def _bold(cell, size: int | None = None) -> None:
    cell.font = Font(bold=True, size=size) if size else Font(bold=True)


def build_sources_uses(ws, *, fill_formulas: bool) -> None:
    """Sources and Uses sheet.

    Row 1: title
    Row 3: Sources header
    Row 4: TLB
    Row 5: Mezzanine
    Row 6: Sponsor Equity (plug)
    Row 7: Total Sources
    Row 9: Uses header
    Row 10: Purchase price (EV)
    Row 11: Transaction fees
    Row 12: Total Uses
    Row 14: Check (sources − uses = 0)
    """
    ws.title = "Sources and Uses"
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 18

    ws.cell(row=1, column=1, value="Sources and Uses — Reed Industries")
    _bold(ws.cell(row=1, column=1), size=13)

    ws.cell(row=3, column=1, value="Sources")
    _bold(ws.cell(row=3, column=1))
    ws.cell(row=4, column=1, value="Term Loan B")
    ws.cell(row=4, column=2, value=TLB0).number_format = "$#,##0"
    ws.cell(row=5, column=1, value="Mezzanine")
    ws.cell(row=5, column=2, value=MEZZ).number_format = "$#,##0"
    ws.cell(row=6, column=1, value="Sponsor Equity (plug)")
    _bold(ws.cell(row=6, column=1))
    ws.cell(row=7, column=1, value="Total Sources")
    _bold(ws.cell(row=7, column=1))

    ws.cell(row=9, column=1, value="Uses")
    _bold(ws.cell(row=9, column=1))
    ws.cell(row=10, column=1, value="Purchase price (EV)")
    ws.cell(row=10, column=2, value=EV).number_format = "$#,##0"
    ws.cell(row=11, column=1, value="Transaction fees")
    ws.cell(row=11, column=2, value=FEES).number_format = "$#,##0"
    ws.cell(row=12, column=1, value="Total Uses")
    _bold(ws.cell(row=12, column=1))

    ws.cell(row=14, column=1, value="Check (Sources − Uses)")

    if fill_formulas:
        # Total Uses first (so Sponsor plug can reference it)
        ws.cell(row=12, column=2, value="=SUM(B10:B11)").number_format = "$#,##0"
        # Sponsor equity = Total Uses − TLB − Mezz
        ws.cell(row=6, column=2, value="=B12-B4-B5").number_format = "$#,##0"
        # Total Sources
        ws.cell(row=7, column=2, value="=SUM(B4:B6)").number_format = "$#,##0"
        # Check
        ws.cell(row=14, column=2, value="=B7-B12").number_format = "$#,##0"
    else:
        for r in (6, 7, 12, 14):
            ws.cell(row=r, column=2).number_format = "$#,##0"


def build_operating_model(ws, *, fill_formulas: bool) -> None:
    """Operating Model sheet.

    Columns: A = label, B..F = Y1..Y5
    Rows:
      1  title
      2  year headers
      3  Revenue
      4  EBITDA
      5  D&A
      6  EBIT
      7  TLB beginning balance
      8  TLB interest
      9  Mezz interest
      10 Pretax income
      11 Tax
      12 Net income
      13 Capex
      14 FCF
      15 Mandatory amort
      16 Cash sweep
      17 TLB end balance
      18 Mezz end balance
    """
    ws.title = "Operating Model"
    ws.column_dimensions["A"].width = 28
    for col in ("B", "C", "D", "E", "F"):
        ws.column_dimensions[col].width = 16

    cols = ["B", "C", "D", "E", "F"]

    ws.cell(row=1, column=1, value="Operating Model — Reed Industries")
    _bold(ws.cell(row=1, column=1), size=13)

    # Year headers
    for i, col in enumerate(cols, 1):
        c = ws.cell(row=2, column=1 + i, value=f"Y{i}")
        _bold(c)
        c.alignment = Alignment(horizontal="center")

    # Labels
    ws.cell(row=3, column=1, value="Revenue")
    ws.cell(row=4, column=1, value="EBITDA")
    ws.cell(row=5, column=1, value="D&A")
    ws.cell(row=6, column=1, value="EBIT")
    ws.cell(row=7, column=1, value="TLB beginning balance")
    ws.cell(row=8, column=1, value="TLB interest")
    ws.cell(row=9, column=1, value="Mezz interest")
    ws.cell(row=10, column=1, value="Pretax income")
    ws.cell(row=11, column=1, value="Tax")
    ws.cell(row=12, column=1, value="Net income")
    ws.cell(row=13, column=1, value="Capex")
    ws.cell(row=14, column=1, value="FCF")
    ws.cell(row=15, column=1, value="Mandatory amort")
    ws.cell(row=16, column=1, value="Cash sweep")
    ws.cell(row=17, column=1, value="TLB end balance")
    _bold(ws.cell(row=17, column=1))
    ws.cell(row=18, column=1, value="Mezz end balance")

    dollar_fmt = "$#,##0"
    for r in range(3, 19):
        for col in cols:
            ws[f"{col}{r}"].number_format = dollar_fmt

    if fill_formulas:
        for i, col in enumerate(cols, 1):
            # Revenue = 250M * 1.06^i
            ws[f"{col}3"] = f"={REV0}*(1+{REV_GROWTH})^{i}"
            # EBITDA = 50M * 1.08^i
            ws[f"{col}4"] = f"={EBITDA0}*(1+{EBITDA_GROWTH})^{i}"
            # D&A flat
            ws[f"{col}5"] = f"={DA_FLAT}"
            # EBIT
            ws[f"{col}6"] = f"={col}4-{col}5"
            # TLB beginning: Y1 references Sources and Uses!B4; Y2+ = prior year end (col-1 row 17)
            if i == 1:
                ws[f"{col}7"] = "='Sources and Uses'!B4"
            else:
                prev_col = cols[i - 2]
                ws[f"{col}7"] = f"={prev_col}17"
            # TLB interest
            ws[f"{col}8"] = f"={TLB_RATE}*{col}7"
            # Mezz interest flat
            ws[f"{col}9"] = f"={MEZZ_INTEREST}"
            # Pretax
            ws[f"{col}10"] = f"={col}6-{col}8-{col}9"
            # Tax
            ws[f"{col}11"] = f"=MAX(0,{col}10)*{TAX}"
            # NI
            ws[f"{col}12"] = f"={col}10-{col}11"
            # Capex
            ws[f"{col}13"] = f"={CAPEX_PCT}*{col}3"
            # FCF
            ws[f"{col}14"] = f"={col}12+{col}5-{col}13"
            # Mandatory amort flat
            ws[f"{col}15"] = f"={MAND_AMORT}"
            # Sweep
            ws[f"{col}16"] = f"=MAX(0,MIN({col}14-{col}15,{col}7-{col}15))"
            # TLB end
            ws[f"{col}17"] = f"={col}7-{col}15-{col}16"
            # Mezz end flat
            ws[f"{col}18"] = f"={MEZZ}"


def build_returns(ws, *, fill_formulas: bool) -> None:
    """Returns sheet.

    Row 1: title
    Row 3: Exit multiple (input)
    Row 4: Hold period (input)
    Row 6: Initial sponsor equity
    Row 7: Exit Y5 EBITDA
    Row 8: Exit EV
    Row 9: TLB at exit
    Row 10: Mezz at exit
    Row 11: Net debt at exit
    Row 12: Exit equity value
    Row 14: MOIC
    Row 15: IRR
    """
    ws.title = "Returns"
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 18

    ws.cell(row=1, column=1, value="Returns — Reed Industries")
    _bold(ws.cell(row=1, column=1), size=13)

    ws.cell(row=3, column=1, value="Exit multiple")
    ws.cell(row=3, column=2, value=EXIT_MULT).number_format = "0.0\"x\""
    ws.cell(row=4, column=1, value="Hold period (years)")
    ws.cell(row=4, column=2, value=HOLD).number_format = "0"

    ws.cell(row=6, column=1, value="Initial sponsor equity")
    _bold(ws.cell(row=6, column=1))
    ws.cell(row=7, column=1, value="Exit Y5 EBITDA")
    ws.cell(row=8, column=1, value="Exit EV")
    _bold(ws.cell(row=8, column=1))
    ws.cell(row=9, column=1, value="TLB at exit")
    ws.cell(row=10, column=1, value="Mezz at exit")
    ws.cell(row=11, column=1, value="Net debt at exit")
    ws.cell(row=12, column=1, value="Exit equity value")
    _bold(ws.cell(row=12, column=1))

    ws.cell(row=14, column=1, value="MOIC")
    _bold(ws.cell(row=14, column=1))
    ws.cell(row=15, column=1, value="IRR")
    _bold(ws.cell(row=15, column=1))

    dollar_fmt = "$#,##0"
    for r in (6, 7, 8, 9, 10, 11, 12):
        ws.cell(row=r, column=2).number_format = dollar_fmt
    ws.cell(row=14, column=2).number_format = "0.00\"x\""
    ws.cell(row=15, column=2).number_format = "0.00%"

    if fill_formulas:
        # Initial sponsor equity from S&U
        ws["B6"] = "='Sources and Uses'!B6"
        # Y5 EBITDA from Operating Model F4
        ws["B7"] = "='Operating Model'!F4"
        # Exit EV = exit_mult × Y5 EBITDA
        ws["B8"] = "=B3*B7"
        # TLB at exit = Operating Model F17
        ws["B9"] = "='Operating Model'!F17"
        # Mezz at exit = Operating Model F18
        ws["B10"] = "='Operating Model'!F18"
        # Net debt = TLB + Mezz
        ws["B11"] = "=B9+B10"
        # Exit equity
        ws["B12"] = "=B8-B11"
        # MOIC
        ws["B14"] = "=B12/B6"
        # IRR
        ws["B15"] = "=(B12/B6)^(1/B4)-1"


# --- files ------------------------------------------------------------------


def build_gold(path: Path) -> None:
    wb = openpyxl.Workbook()
    build_sources_uses(wb.active, fill_formulas=True)
    ws2 = wb.create_sheet(title="Operating Model")
    build_operating_model(ws2, fill_formulas=True)
    ws3 = wb.create_sheet(title="Returns")
    build_returns(ws3, fill_formulas=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    recalc_xlsx(path)


def build_stub(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ReadMe"
    ws["A1"] = "Builder: create 'Sources and Uses', 'Operating Model', and 'Returns' sheets."
    ws["A2"] = "See brief.md, inputs/deal_terms.md, inputs/model_format.md."
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def build_results_json(path: Path) -> None:
    """Map output keys to their Sheet!Cell addresses."""
    data = {
        # Sources & Uses
        "total_uses": "Sources and Uses!B12",
        "total_sources": "Sources and Uses!B7",
        "sponsor_initial_equity": "Sources and Uses!B6",
        # Operating Model
        "y1_ebitda": "Operating Model!B4",
        "y5_ebitda": "Operating Model!F4",
        "y5_fcf": "Operating Model!F14",
        "tlb_y5_end_balance": "Operating Model!F17",
        # Returns
        "exit_ev": "Returns!B8",
        "net_debt_at_exit": "Returns!B11",
        "exit_equity_value": "Returns!B12",
        "sponsor_moic": "Returns!B14",
        "sponsor_irr": "Returns!B15",
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, indent=2))


def main() -> None:
    build_stub(TASK_DIR / "stubs" / "starter.xlsx")
    build_gold(TASK_DIR / "gold" / "model.xlsx")
    build_results_json(TASK_DIR / "gold" / "results.json")
    print("t2_lbo_mini artifacts written.")
    print(f"  total_uses       = ${GOLD['total_uses']:>16,.2f}")
    print(f"  sponsor_equity   = ${GOLD['sponsor_eq']:>16,.2f}")
    print(f"  Y1 EBITDA        = ${GOLD['years'][0]['ebitda']:>16,.2f}")
    print(f"  Y5 EBITDA        = ${GOLD['ebitda_y5']:>16,.2f}")
    print(f"  Y5 FCF           = ${GOLD['y5_fcf']:>16,.2f}")
    print(f"  TLB Y5 end       = ${GOLD['tlb_y5_end']:>16,.2f}")
    print(f"  Exit EV          = ${GOLD['exit_ev']:>16,.2f}")
    print(f"  Net debt exit    = ${GOLD['net_debt']:>16,.2f}")
    print(f"  Exit equity      = ${GOLD['exit_equity']:>16,.2f}")
    print(f"  MOIC             = {GOLD['moic']:.4f}x")
    print(f"  IRR              = {GOLD['irr']*100:.2f}%")


if __name__ == "__main__":
    main()
