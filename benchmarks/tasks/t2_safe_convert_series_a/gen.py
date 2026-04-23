"""Generator for t2_safe_convert_series_a.

Three-sheet model:
  1. Cap Table   — copy of the input pre-round cap table (verbatim)
  2. SAFE Convert — per-SAFE conversion price and share count
  3. Series A    — headline price, SAFE block, new investors, pool top-up

Math (non-circular by construction):
  - pre_FD = 10_500_000  (from Cap Table)
  - headline_price = pre_money / pre_FD
  - for each SAFE:
      cap_price = cap / pre_FD
      discount_price = (1 - discount) * headline_price
      conv_price = min(cap_price, discount_price)
      safe_shares = investment / conv_price
  - new_investor_shares = round_size / headline_price
  - pool top-up solves: (unalloc + top_up) = 10% * (pre_FD + safe_total + new_inv + top_up)
    => top_up = (pool_pct * (pre_FD + safe_total + new_inv) - unalloc) / (1 - pool_pct)
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from benchmarks.recalc import recalc_xlsx

TASK_DIR = Path(__file__).resolve().parent

# --- parameters -------------------------------------------------------------
PRE_FD = 10_500_000
PRE_UNALLOC = 500_000

SAFE1 = {"name": "Harbor Fund", "investment": 2_000_000, "cap": 20_000_000, "discount": 0.0}
SAFE2 = {"name": "Archer Angels", "investment": 1_500_000, "cap": 25_000_000, "discount": 0.20}

PRE_MONEY = 40_000_000
ROUND_SIZE = 15_000_000
POOL_PCT = 0.10

# Cap table rows: (holder, class, shares)
CAP_ROWS = [
    ("Founder A", "Common", 4_000_000),
    ("Founder B", "Common", 3_000_000),
    ("Seed Investor 1", "Preferred Seed", 1_500_000),
    ("Seed Investor 2", "Preferred Seed", 1_000_000),
    ("Option Pool (granted)", "Options (granted)", 500_000),
    ("Option Pool (unallocated)", "Options (unallocated)", PRE_UNALLOC),
]


# --- gold math --------------------------------------------------------------


def compute_gold() -> dict:
    headline = PRE_MONEY / PRE_FD
    def safe_calc(s):
        cap_price = s["cap"] / PRE_FD
        disc_price = (1 - s["discount"]) * headline
        conv = min(cap_price, disc_price)
        shares = s["investment"] / conv
        return {"cap_price": cap_price, "disc_price": disc_price,
                "conv_price": conv, "shares": shares}
    s1 = safe_calc(SAFE1)
    s2 = safe_calc(SAFE2)
    safe_total_shares = s1["shares"] + s2["shares"]
    new_inv_shares = ROUND_SIZE / headline
    top_up = (POOL_PCT * (PRE_FD + safe_total_shares + new_inv_shares) - PRE_UNALLOC) / (1 - POOL_PCT)
    post_fd = PRE_FD + safe_total_shares + new_inv_shares + top_up
    return {
        "headline": headline,
        "safe1": s1, "safe2": s2,
        "safe_total": safe_total_shares,
        "new_inv": new_inv_shares,
        "pool_topup": top_up,
        "post_fd": post_fd,
    }


GOLD = compute_gold()


# --- builders ---------------------------------------------------------------


def _hdr(cell) -> None:
    cell.font = Font(bold=True)
    cell.fill = PatternFill("solid", fgColor="DDEEFF")


def build_cap_table(ws, *, fill_formulas: bool) -> None:
    ws.title = "Cap Table"
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 10
    headers = ["Holder", "Class", "Shares", "% FD"]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8E8E8")
    for r, (h, cls, sh) in enumerate(CAP_ROWS, 2):
        ws.cell(row=r, column=1, value=h)
        ws.cell(row=r, column=2, value=cls)
        ws.cell(row=r, column=3, value=sh).number_format = "#,##0"
        if fill_formulas:
            ws.cell(row=r, column=4, value=f"=C{r}/$C${len(CAP_ROWS)+2}").number_format = "0.00%"
        else:
            ws.cell(row=r, column=4).number_format = "0.00%"
    total_row = len(CAP_ROWS) + 2
    ws.cell(row=total_row, column=1, value="Total FD").font = Font(bold=True)
    if fill_formulas:
        ws.cell(row=total_row, column=3, value=f"=SUM(C2:C{total_row-1})").number_format = "#,##0"
        ws.cell(row=total_row, column=4, value=f"=SUM(D2:D{total_row-1})").number_format = "0.00%"


def build_safe_convert(ws, *, fill_formulas: bool) -> None:
    ws.title = "SAFE Convert"
    ws.column_dimensions["A"].width = 32
    for c in ("B", "C"):
        ws.column_dimensions[c].width = 18

    ws.cell(row=1, column=1, value="SAFE Conversion").font = Font(bold=True, size=13)
    ws.cell(row=3, column=2, value="Harbor Fund").font = Font(bold=True)
    ws.cell(row=3, column=3, value="Archer Angels").font = Font(bold=True)

    # Row map:
    #  4 Investment
    #  5 Valuation cap
    #  6 Discount
    #  7 Pre-round FD (referenced)
    #  8 Headline price
    #  9 Cap price
    # 10 Discount price
    # 11 Conversion price
    # 12 Shares issued
    rows = [
        ("Investment", SAFE1["investment"], SAFE2["investment"]),
        ("Valuation cap", SAFE1["cap"], SAFE2["cap"]),
        ("Discount", SAFE1["discount"], SAFE2["discount"]),
    ]
    for i, (label, v1, v2) in enumerate(rows, 4):
        ws.cell(row=i, column=1, value=label)
        fmt = "0.0%" if label == "Discount" else "$#,##0"
        ws.cell(row=i, column=2, value=v1).number_format = fmt
        ws.cell(row=i, column=3, value=v2).number_format = fmt

    ws.cell(row=7, column=1, value="Pre-round FD (from Cap Table)")
    if fill_formulas:
        ws.cell(row=7, column=2, value="='Cap Table'!C8").number_format = "#,##0"
        ws.cell(row=7, column=3, value="='Cap Table'!C8").number_format = "#,##0"
    else:
        ws.cell(row=7, column=2).number_format = "#,##0"
        ws.cell(row=7, column=3).number_format = "#,##0"

    ws.cell(row=8, column=1, value="Headline price (from Series A)")
    if fill_formulas:
        ws.cell(row=8, column=2, value="='Series A'!B5").number_format = "$0.0000"
        ws.cell(row=8, column=3, value="='Series A'!B5").number_format = "$0.0000"
    else:
        ws.cell(row=8, column=2).number_format = "$0.0000"
        ws.cell(row=8, column=3).number_format = "$0.0000"

    ws.cell(row=9, column=1, value="Cap price")
    ws.cell(row=10, column=1, value="Discount price")
    ws.cell(row=11, column=1, value="Conversion price").font = Font(bold=True)
    ws.cell(row=12, column=1, value="Shares issued").font = Font(bold=True)

    if fill_formulas:
        ws.cell(row=9, column=2, value="=B5/B7").number_format = "$0.0000"
        ws.cell(row=9, column=3, value="=C5/C7").number_format = "$0.0000"
        ws.cell(row=10, column=2, value="=(1-B6)*B8").number_format = "$0.0000"
        ws.cell(row=10, column=3, value="=(1-C6)*C8").number_format = "$0.0000"
        ws.cell(row=11, column=2, value="=MIN(B9,B10)").number_format = "$0.0000"
        ws.cell(row=11, column=3, value="=MIN(C9,C10)").number_format = "$0.0000"
        ws.cell(row=12, column=2, value="=B4/B11").number_format = "#,##0"
        ws.cell(row=12, column=3, value="=C4/C11").number_format = "#,##0"
    else:
        for rr in (9, 10, 11):
            ws.cell(row=rr, column=2).number_format = "$0.0000"
            ws.cell(row=rr, column=3).number_format = "$0.0000"
        for cc in ("B", "C"):
            ws[f"{cc}12"].number_format = "#,##0"

    ws.cell(row=14, column=1, value="Total SAFE shares").font = Font(bold=True)
    if fill_formulas:
        ws.cell(row=14, column=2, value="=B12+C12").number_format = "#,##0"


def build_series_a(ws, *, fill_formulas: bool) -> None:
    ws.title = "Series A"
    ws.column_dimensions["A"].width = 38
    for c in ("B", "C", "D"):
        ws.column_dimensions[c].width = 16

    ws.cell(row=1, column=1, value="Series A — round mechanics").font = Font(bold=True, size=13)

    # Section: round inputs
    ws.cell(row=3, column=1, value="Pre-money valuation")
    ws.cell(row=3, column=2, value=PRE_MONEY).number_format = "$#,##0"
    ws.cell(row=4, column=1, value="Pre-round FD (from Cap Table)")
    if fill_formulas:
        ws.cell(row=4, column=2, value="='Cap Table'!C8").number_format = "#,##0"
    else:
        ws.cell(row=4, column=2).number_format = "#,##0"
    ws.cell(row=5, column=1, value="Headline price per share").font = Font(bold=True)
    if fill_formulas:
        ws.cell(row=5, column=2, value="=B3/B4").number_format = "$0.0000"
    else:
        ws.cell(row=5, column=2).number_format = "$0.0000"

    ws.cell(row=6, column=1, value="Round size (new money)")
    ws.cell(row=6, column=2, value=ROUND_SIZE).number_format = "$#,##0"
    ws.cell(row=7, column=1, value="Option pool target (post-money)")
    ws.cell(row=7, column=2, value=POOL_PCT).number_format = "0.0%"
    ws.cell(row=8, column=1, value="Pre-round unallocated pool")
    ws.cell(row=8, column=2, value=PRE_UNALLOC).number_format = "#,##0"

    # Section: SAFE + new investor + pool
    ws.cell(row=11, column=1, value="SAFE shares (from SAFE Convert)")
    ws.cell(row=12, column=1, value="New investor shares")
    ws.cell(row=13, column=1, value="Option pool top-up").font = Font(bold=True)
    ws.cell(row=14, column=1, value="Post-money fully diluted shares").font = Font(bold=True)
    if fill_formulas:
        ws.cell(row=11, column=2, value="='SAFE Convert'!B14").number_format = "#,##0"
        ws.cell(row=12, column=2, value="=B6/B5").number_format = "#,##0"
        ws.cell(row=13, column=2, value="=(B7*(B4+B11+B12)-B8)/(1-B7)").number_format = "#,##0"
        ws.cell(row=14, column=2, value="=B4+B11+B12+B13").number_format = "#,##0"
    else:
        for rr in (11, 12, 13, 14):
            ws.cell(row=rr, column=2).number_format = "#,##0"


def build_gold(path: Path) -> None:
    wb = openpyxl.Workbook()
    # Replace default sheet by building first sheet then adding others
    build_cap_table(wb.active, fill_formulas=True)
    ws2 = wb.create_sheet(title="SAFE Convert")
    build_safe_convert(ws2, fill_formulas=True)
    ws3 = wb.create_sheet(title="Series A")
    build_series_a(ws3, fill_formulas=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    recalc_xlsx(path)


def build_input_cap_table(path: Path) -> None:
    wb = openpyxl.Workbook()
    build_cap_table(wb.active, fill_formulas=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    recalc_xlsx(path)


def build_stub(path: Path) -> None:
    wb = openpyxl.Workbook()
    wb.active.title = "ReadMe"
    wb.active["A1"] = "Builder: create Cap Table, SAFE Convert, and Series A sheets."
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def build_results_json(path: Path) -> None:
    """Gold results.json: semantic keys → Sheet!Cell addresses.

    Matches the layout produced by build_cap_table / build_safe_convert /
    build_series_a with fill_formulas=True.
    """
    import json
    data = {
        # Cap Table: 6 holder rows + total row at row 8
        "cap_total_fd": "Cap Table!C8",
        # SAFE Convert: row 9 cap, 10 discount, 11 conv, 12 shares; B=SAFE1, C=SAFE2; row 14 total
        "safe1_cap_price": "SAFE Convert!B9",
        "safe1_discount_price": "SAFE Convert!B10",
        "safe1_conversion_price": "SAFE Convert!B11",
        "safe1_shares_issued": "SAFE Convert!B12",
        "safe2_cap_price": "SAFE Convert!C9",
        "safe2_discount_price": "SAFE Convert!C10",
        "safe2_conversion_price": "SAFE Convert!C11",
        "safe2_shares_issued": "SAFE Convert!C12",
        "total_safe_shares": "SAFE Convert!B14",
        # Series A: row 5 headline, 11 safe shares, 12 new inv, 13 topup, 14 post FD
        "series_a_headline_price": "Series A!B5",
        "series_a_new_investor_shares": "Series A!B12",
        "series_a_pool_topup": "Series A!B13",
        "series_a_post_fd": "Series A!B14",
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, indent=2))


def main() -> None:
    build_input_cap_table(TASK_DIR / "inputs" / "pre_cap_table.xlsx")
    build_stub(TASK_DIR / "stubs" / "starter.xlsx")
    build_gold(TASK_DIR / "gold" / "model.xlsx")
    build_results_json(TASK_DIR / "gold" / "results.json")
    print("t2 artifacts written.")
    print(f"  headline = ${GOLD['headline']:.4f}")
    print(f"  SAFE1 conv = ${GOLD['safe1']['conv_price']:.4f}, shares = {GOLD['safe1']['shares']:,.2f}")
    print(f"  SAFE2 conv = ${GOLD['safe2']['conv_price']:.4f}, shares = {GOLD['safe2']['shares']:,.2f}")
    print(f"  new investor shares = {GOLD['new_inv']:,.2f}")
    print(f"  pool top-up = {GOLD['pool_topup']:,.2f}")
    print(f"  post-money FD = {GOLD['post_fd']:,.2f}")


if __name__ == "__main__":
    main()
