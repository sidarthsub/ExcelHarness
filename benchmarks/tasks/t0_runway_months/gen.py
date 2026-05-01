"""Generator for t0_runway_months.

Deterministically produces:
  stubs/starter.xlsx    - workbook with Calc sheet inputs filled, answers blank
  gold/model.xlsx       - same as stub but with formula-driven answers
  gold/results.json     - maps semantic keys to Calc cell addresses
  gold/grading.yaml     - rubric (hardcoded, with the computed expected values)

Math: cash runway with a stepped monthly burn.
  Starting cash:          $20,000,000
  Monthly burn (mo 1-6):   $1,000,000
  Monthly burn (mo 7+):    $1,500,000
  Phase 1 length:          6 months

  cash_end_phase_1 = starting_cash - phase_1_len * burn_1
                   = 20M - 6 * 1M = 14M
  phase_2_runway   = cash_end_phase_1 / burn_2
                   = 14M / 1.5M = 9.3333...
  total_runway     = phase_1_len + phase_2_runway
                   = 6 + 9.3333... = 15.3333...
"""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

from recalc import recalc_xlsx

TASK_DIR = Path(__file__).resolve().parent

# --- parameters -------------------------------------------------------------
STARTING_CASH = 20_000_000
BURN_PHASE_1 = 1_000_000
BURN_PHASE_2 = 1_500_000
PHASE_1_LEN = 6

# Expected values, computed at gen time so grading.yaml stays in sync.
GOLD_CASH_END_PHASE_1 = STARTING_CASH - PHASE_1_LEN * BURN_PHASE_1
GOLD_PHASE_2_RUNWAY = GOLD_CASH_END_PHASE_1 / BURN_PHASE_2
GOLD_TOTAL_RUNWAY = PHASE_1_LEN + GOLD_PHASE_2_RUNWAY


# --- builders ---------------------------------------------------------------


def build_starter(path: Path, *, fill_answer: bool) -> None:
    """Build a workbook with a single Calc sheet.

    Layout:
      A1  "Runway calculator"
      A3  "Starting cash"              B3  20,000,000
      A4  "Burn mo 1-6"                B4   1,000,000
      A5  "Burn mo 7+"                 B5   1,500,000
      A6  "Phase 1 length"             B6   6
      A9  "Cash at end of phase 1"     B9  <formula if fill_answer>
      A10 "Phase 2 months of runway"   B10 <formula if fill_answer>
      A11 "Total runway (months)"      B11 <formula if fill_answer>
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Calc"

    hdr = Font(bold=True)
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 18

    ws.cell(row=1, column=1, value="Runway calculator").font = hdr

    ws.cell(row=3, column=1, value="Starting cash")
    ws.cell(row=3, column=2, value=STARTING_CASH).number_format = "$#,##0"

    ws.cell(row=4, column=1, value="Burn mo 1-6")
    ws.cell(row=4, column=2, value=BURN_PHASE_1).number_format = "$#,##0"

    ws.cell(row=5, column=1, value="Burn mo 7+")
    ws.cell(row=5, column=2, value=BURN_PHASE_2).number_format = "$#,##0"

    ws.cell(row=6, column=1, value="Phase 1 length")
    ws.cell(row=6, column=2, value=PHASE_1_LEN).number_format = "0"

    ws.cell(row=9, column=1, value="Cash at end of phase 1")
    if fill_answer:
        ws.cell(row=9, column=2, value="=B3-B6*B4")
    ws.cell(row=9, column=2).number_format = "$#,##0.00"

    ws.cell(row=10, column=1, value="Phase 2 months of runway")
    if fill_answer:
        ws.cell(row=10, column=2, value="=B9/B5")
    ws.cell(row=10, column=2).number_format = "0.0000"

    ws.cell(row=11, column=1, value="Total runway (months)").font = hdr
    if fill_answer:
        ws.cell(row=11, column=2, value="=B6+B10")
    ws.cell(row=11, column=2).number_format = "0.0000"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def build_results_json(path: Path) -> None:
    data = {
        "cash_end_phase_1": "Calc!B9",
        "phase_2_runway": "Calc!B10",
        "total_runway_months": "Calc!B11",
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, indent=2))


def build_grading_yaml(path: Path) -> None:
    """Emit grading.yaml with the computed expected values baked in."""
    content = f"""recalc: true
checks:
  - id: calc_sheet_exists
    type: sheet_exists
    sheet: Calc
    weight: 0.3

  - id: cash_end_phase_1_value
    type: output_key
    key: cash_end_phase_1
    expected: {GOLD_CASH_END_PHASE_1!r}
    tolerance_rel: 1.0e-6
    require_formula: true
    weight: 0.7
    description: "Cash remaining after phase 1 must match and be formula-driven."

  - id: phase_2_runway_value
    type: output_key
    key: phase_2_runway
    expected: {GOLD_PHASE_2_RUNWAY!r}
    tolerance_rel: 1.0e-6
    require_formula: true
    weight: 0.8
    description: "Phase 2 months of runway must match and be formula-driven."

  - id: total_runway_months_value
    type: output_key
    key: total_runway_months
    expected: {GOLD_TOTAL_RUNWAY!r}
    tolerance_rel: 1.0e-6
    require_formula: true
    weight: 1.2
    description: "Total runway (headline answer) must match and be formula-driven."
"""
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content)


def main() -> None:
    build_starter(TASK_DIR / "stubs" / "starter.xlsx", fill_answer=False)

    gold = TASK_DIR / "gold" / "model.xlsx"
    build_starter(gold, fill_answer=True)
    recalc_xlsx(gold)

    build_results_json(TASK_DIR / "gold" / "results.json")
    build_grading_yaml(TASK_DIR / "gold" / "grading.yaml")

    print(f"GOLD_CASH_END_PHASE_1 = {GOLD_CASH_END_PHASE_1:,.4f}")
    print(f"GOLD_PHASE_2_RUNWAY   = {GOLD_PHASE_2_RUNWAY:,.6f}")
    print(f"GOLD_TOTAL_RUNWAY     = {GOLD_TOTAL_RUNWAY:,.6f}")
    print("t0_runway_months artifacts written.")


if __name__ == "__main__":
    main()
