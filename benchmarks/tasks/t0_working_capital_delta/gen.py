"""Generator for t0_working_capital_delta.

Deterministically produces:
  stubs/starter.xlsx       — workbook with inputs filled, answer cells blank
  gold/model.xlsx          — same as stub but with correct formulas
  gold/results.json        — semantic keys → Sheet!Cell addresses

Math: YoY working capital cascade (365-day, source-positive convention).

  AR_t  = Revenue_t × DSO / 365
  AP_t  = COGS_t    × DPO / 365    (COGS_t = Revenue_t × cogs_pct)
  Inv_t = COGS_t    × DIO / 365
  ΔAR  = AR_Y1  − AR_Y0
  ΔAP  = AP_Y1  − AP_Y0
  ΔInv = Inv_Y1 − Inv_Y0
  Net WC cash impact = −ΔAR + ΔAP − ΔInv

With the given parameters:
  Revenue Y0 = $80M, Y1 = $100M, COGS % = 0.30, DSO = 45, DPO = 60, DIO = 30.
"""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

from benchmarks.recalc import recalc_xlsx

TASK_DIR = Path(__file__).resolve().parent

# --- parameters -------------------------------------------------------------
REV_Y0 = 80_000_000
REV_Y1 = 100_000_000
COGS_PCT = 0.30
DSO = 45
DPO = 60
DIO = 30

# --- gold values (for logging) ----------------------------------------------
GOLD_AR_Y0 = REV_Y0 * DSO / 365
GOLD_AR_Y1 = REV_Y1 * DSO / 365
GOLD_DELTA_AR = GOLD_AR_Y1 - GOLD_AR_Y0

GOLD_AP_Y0 = REV_Y0 * COGS_PCT * DPO / 365
GOLD_AP_Y1 = REV_Y1 * COGS_PCT * DPO / 365
GOLD_DELTA_AP = GOLD_AP_Y1 - GOLD_AP_Y0

GOLD_INV_Y0 = REV_Y0 * COGS_PCT * DIO / 365
GOLD_INV_Y1 = REV_Y1 * COGS_PCT * DIO / 365
GOLD_DELTA_INV = GOLD_INV_Y1 - GOLD_INV_Y0

GOLD_NET_WC = -GOLD_DELTA_AR + GOLD_DELTA_AP - GOLD_DELTA_INV


# --- builders ---------------------------------------------------------------


def build_starter(path: Path, *, fill_formulas: bool) -> None:
    """Build the Calc sheet. If fill_formulas, populate the answer cells
    (B11:B20) with formulas referencing B3:B8. Otherwise leave blank.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Calc"

    hdr = Font(bold=True)
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 20

    ws.cell(row=1, column=1, value="Working capital cascade").font = hdr

    ws.cell(row=3, column=1, value="Revenue Y0")
    ws.cell(row=3, column=2, value=REV_Y0).number_format = "#,##0"

    ws.cell(row=4, column=1, value="Revenue Y1")
    ws.cell(row=4, column=2, value=REV_Y1).number_format = "#,##0"

    ws.cell(row=5, column=1, value="COGS % of rev")
    ws.cell(row=5, column=2, value=COGS_PCT).number_format = "0.00%"

    ws.cell(row=6, column=1, value="DSO")
    ws.cell(row=6, column=2, value=DSO).number_format = "0"

    ws.cell(row=7, column=1, value="DPO")
    ws.cell(row=7, column=2, value=DPO).number_format = "0"

    ws.cell(row=8, column=1, value="DIO")
    ws.cell(row=8, column=2, value=DIO).number_format = "0"

    ws.cell(row=10, column=1, value="Answers:").font = hdr

    ws.cell(row=11, column=1, value="AR Y0")
    ws.cell(row=12, column=1, value="AR Y1")
    ws.cell(row=13, column=1, value="ΔAR")
    ws.cell(row=14, column=1, value="AP Y0")
    ws.cell(row=15, column=1, value="AP Y1")
    ws.cell(row=16, column=1, value="ΔAP")
    ws.cell(row=17, column=1, value="Inv Y0")
    ws.cell(row=18, column=1, value="Inv Y1")
    ws.cell(row=19, column=1, value="ΔInv")
    ws.cell(row=20, column=1, value="Net WC cash impact")

    if fill_formulas:
        ws.cell(row=11, column=2, value="=B3*B6/365")
        ws.cell(row=12, column=2, value="=B4*B6/365")
        ws.cell(row=13, column=2, value="=B12-B11")
        ws.cell(row=14, column=2, value="=B3*B5*B7/365")
        ws.cell(row=15, column=2, value="=B4*B5*B7/365")
        ws.cell(row=16, column=2, value="=B15-B14")
        ws.cell(row=17, column=2, value="=B3*B5*B8/365")
        ws.cell(row=18, column=2, value="=B4*B5*B8/365")
        ws.cell(row=19, column=2, value="=B18-B17")
        ws.cell(row=20, column=2, value="=-B13+B16-B19")

    for r in range(11, 21):
        ws.cell(row=r, column=2).number_format = "#,##0.00"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def build_results_json(path: Path) -> None:
    data = {
        "ar_y0": "Calc!B11",
        "ar_y1": "Calc!B12",
        "delta_ar": "Calc!B13",
        "ap_y0": "Calc!B14",
        "ap_y1": "Calc!B15",
        "delta_ap": "Calc!B16",
        "inv_y0": "Calc!B17",
        "inv_y1": "Calc!B18",
        "delta_inv": "Calc!B19",
        "net_wc_cash_impact": "Calc!B20",
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, indent=2))


def main() -> None:
    build_starter(TASK_DIR / "stubs" / "starter.xlsx", fill_formulas=False)
    gold = TASK_DIR / "gold" / "model.xlsx"
    build_starter(gold, fill_formulas=True)
    recalc_xlsx(gold)
    build_results_json(TASK_DIR / "gold" / "results.json")
    print("t0_working_capital_delta artifacts written.")
    print(f"  ar_y0              = {GOLD_AR_Y0:,.3f}")
    print(f"  ar_y1              = {GOLD_AR_Y1:,.3f}")
    print(f"  delta_ar           = {GOLD_DELTA_AR:,.3f}")
    print(f"  ap_y0              = {GOLD_AP_Y0:,.3f}")
    print(f"  ap_y1              = {GOLD_AP_Y1:,.3f}")
    print(f"  delta_ap           = {GOLD_DELTA_AP:,.3f}")
    print(f"  inv_y0             = {GOLD_INV_Y0:,.3f}")
    print(f"  inv_y1             = {GOLD_INV_Y1:,.3f}")
    print(f"  delta_inv          = {GOLD_DELTA_INV:,.3f}")
    print(f"  net_wc_cash_impact = {GOLD_NET_WC:,.3f}")


if __name__ == "__main__":
    main()
