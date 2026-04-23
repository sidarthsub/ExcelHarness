"""Generator for t1_revenue_build.

Task: build a monthly revenue build on a single sheet named "Revenue Build".
Monthly units = annual_units × seasonality; monthly revenue = units × price.
Totals row sums units and revenue; weighted-avg price = total_rev / total_units.

The stub is a near-blank workbook. The gold is the canonical Revenue Build
sheet with formula-driven monthly and total rows, plus a small parameters
area (B18: annual_units, B19: price).
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

from benchmarks.recalc import recalc_xlsx

TASK_DIR = Path(__file__).resolve().parent

# --- parameters -------------------------------------------------------------
ANNUAL_UNITS = 120_000
BASE_PRICE = 50

MONTHS = [
    ("Jan", 0.05),
    ("Feb", 0.05),
    ("Mar", 0.07),
    ("Apr", 0.08),
    ("May", 0.08),
    ("Jun", 0.09),
    ("Jul", 0.10),
    ("Aug", 0.10),
    ("Sep", 0.09),
    ("Oct", 0.10),
    ("Nov", 0.11),
    ("Dec", 0.08),
]

SHEET_NAME = "Revenue Build"


def _hdr(cell) -> None:
    cell.font = Font(bold=True)
    cell.fill = PatternFill("solid", fgColor="DDEEFF")


def build_revenue_sheet(ws, *, fill_formulas: bool) -> None:
    ws.title = SHEET_NAME
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 16

    title = ws.cell(row=1, column=1, value="Revenue Build")
    title.font = Font(bold=True, size=13)

    # Row 3: headers
    headers = ["Month", "Seasonality %", "Units", "Price", "Revenue"]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=3, column=i, value=h)
        _hdr(c)

    # Rows 4-15: Jan-Dec
    first_row = 4
    for i, (name, frac) in enumerate(MONTHS):
        r = first_row + i
        ws.cell(row=r, column=1, value=name)
        ws.cell(row=r, column=2, value=frac).number_format = "0.00%"
        if fill_formulas:
            ws.cell(row=r, column=3, value=f"=$B$18*B{r}").number_format = "#,##0"
            ws.cell(row=r, column=4, value="=$B$19").number_format = "$#,##0.00"
            ws.cell(row=r, column=5, value=f"=C{r}*D{r}").number_format = "$#,##0"
        else:
            ws.cell(row=r, column=3).number_format = "#,##0"
            ws.cell(row=r, column=4).number_format = "$#,##0.00"
            ws.cell(row=r, column=5).number_format = "$#,##0"

    # Row 16: Totals
    last_row = first_row + len(MONTHS) - 1  # 15
    tot_row = last_row + 1  # 16
    ws.cell(row=tot_row, column=1, value="Total").font = Font(bold=True)
    if fill_formulas:
        ws.cell(row=tot_row, column=2, value=f"=SUM(B{first_row}:B{last_row})").number_format = "0.00%"
        ws.cell(row=tot_row, column=3, value=f"=SUM(C{first_row}:C{last_row})").number_format = "#,##0"
        ws.cell(row=tot_row, column=5, value=f"=SUM(E{first_row}:E{last_row})").number_format = "$#,##0"
        # Weighted-avg price depends on total rev + total units already set
        ws.cell(row=tot_row, column=4, value=f"=E{tot_row}/C{tot_row}").number_format = "$#,##0.00"
    else:
        ws.cell(row=tot_row, column=2).number_format = "0.00%"
        ws.cell(row=tot_row, column=3).number_format = "#,##0"
        ws.cell(row=tot_row, column=4).number_format = "$#,##0.00"
        ws.cell(row=tot_row, column=5).number_format = "$#,##0"

    # Parameters area (rows 18-19, col A label / col B value)
    params_header = ws.cell(row=17, column=1, value="Parameters")
    _hdr(params_header)
    ws.cell(row=18, column=1, value="Annual units")
    ws.cell(row=18, column=2, value=ANNUAL_UNITS).number_format = "#,##0"
    ws.cell(row=19, column=1, value="Price")
    ws.cell(row=19, column=2, value=BASE_PRICE).number_format = "$#,##0.00"


def build_stub(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ReadMe"
    ws["A1"] = "Build Revenue Build sheet per brief and inputs/assumptions.md."
    ws["A1"].font = Font(bold=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def build_gold(path: Path) -> None:
    wb = openpyxl.Workbook()
    build_revenue_sheet(wb.active, fill_formulas=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    recalc_xlsx(path)


def build_results_json(path: Path) -> None:
    """Gold results.json: each semantic key mapped to its Sheet!Cell address.

    Layout produced by build_revenue_sheet(fill_formulas=True):
      B16  Sum of seasonality fractions (= 1.0)
      C16  Total annual units (= 120,000)
      E16  Total annual revenue (= 6,000,000)
      D16  Weighted avg price (= 50)
      C10  Jul units (row 4=Jan ... row 10=Jul)
      E10  Jul revenue
      E14  Nov revenue (row 4=Jan ... row 14=Nov)
    """
    import json
    sheet = SHEET_NAME
    # Months start at row 4: Jan=4, Feb=5, ..., Jul=10, ..., Nov=14, Dec=15
    month_row = {name: 4 + i for i, (name, _) in enumerate(MONTHS)}
    data = {
        "seasonality_sum": f"{sheet}!B16",
        "total_annual_units": f"{sheet}!C16",
        "total_annual_revenue": f"{sheet}!E16",
        "jul_units": f"{sheet}!C{month_row['Jul']}",
        "jul_revenue": f"{sheet}!E{month_row['Jul']}",
        "nov_revenue": f"{sheet}!E{month_row['Nov']}",
        "weighted_avg_price": f"{sheet}!D16",
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, indent=2))


def main() -> None:
    build_stub(TASK_DIR / "stubs" / "starter.xlsx")
    build_gold(TASK_DIR / "gold" / "model.xlsx")
    build_results_json(TASK_DIR / "gold" / "results.json")
    print("t1_revenue_build artifacts written.")


if __name__ == "__main__":
    main()
