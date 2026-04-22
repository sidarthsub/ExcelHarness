"""Dump a live workbook snapshot to per-sheet JSON for the evaluator.

Reads the snapshot xlsx with openpyxl twice (formulas + cached values) and
emits one JSON file per sheet in the format the evaluator expects:

    {
        "sheet": <name>,
        "address": "A1:J23",
        "dimensions": "A1:J23",
        "values":       [[...], ...],
        "formulas":     [[...], ...],
        "numberFormat": [[...], ...],
    }

Replaces the Office.js dumpSheet path which was bottlenecked by live-Excel
recalc + RPC marshalling (30s timeouts on heavy sheets).
"""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter


def _used_bounds(ws) -> tuple[int, int] | None:
    """Return (max_row, max_col) trimmed to the last row/col that has a value.

    openpyxl's ws.max_row/max_column can over-report when cells have only
    formatting. We walk backwards to find the true data bounds.
    """
    max_col = ws.max_column or 0
    max_row = ws.max_row or 0
    if not max_col or not max_row:
        return None

    actual_max_row = 0
    for r in range(max_row, 0, -1):
        if any(ws.cell(r, c).value is not None for c in range(1, max_col + 1)):
            actual_max_row = r
            break
    if not actual_max_row:
        return None

    actual_max_col = 0
    for c in range(max_col, 0, -1):
        if any(ws.cell(r, c).value is not None for r in range(1, actual_max_row + 1)):
            actual_max_col = c
            break
    if not actual_max_col:
        return None
    return actual_max_row, actual_max_col


def dump_xlsx_to_json(xlsx_path: Path, out_dir: Path) -> list[Path]:
    """Dump every sheet in xlsx_path to out_dir/<Safe_Sheet_Name>.json.

    Returns the list of written paths.
    """
    xlsx_path = Path(xlsx_path)
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    wb_formulas = openpyxl.load_workbook(xlsx_path, data_only=False)
    wb_values = openpyxl.load_workbook(xlsx_path, data_only=True)

    written: list[Path] = []
    try:
        for ws in wb_formulas.worksheets:
            vs = wb_values[ws.title]
            bounds = _used_bounds(ws)

            if bounds is None:
                address = "A1:A1"
                values = [[None]]
                formulas = [[None]]
                number_formats = [["General"]]
                rows, cols = 1, 1
            else:
                rows, cols = bounds
                address = f"A1:{get_column_letter(cols)}{rows}"
                values = [[None] * cols for _ in range(rows)]
                formulas = [[None] * cols for _ in range(rows)]
                number_formats = [["General"] * cols for _ in range(rows)]
                for r in range(1, rows + 1):
                    for c in range(1, cols + 1):
                        f_cell = ws.cell(r, c)
                        v_cell = vs.cell(r, c)
                        raw = f_cell.value
                        is_formula = isinstance(raw, str) and raw.startswith("=")
                        if is_formula:
                            formulas[r - 1][c - 1] = raw
                            values[r - 1][c - 1] = v_cell.value
                        else:
                            formulas[r - 1][c - 1] = raw
                            values[r - 1][c - 1] = raw
                        number_formats[r - 1][c - 1] = f_cell.number_format or "General"

            payload = {
                "sheet": ws.title,
                "address": address,
                "dimensions": address,
                "values": values,
                "formulas": formulas,
                "numberFormat": number_formats,
            }
            safe_name = ws.title.replace("/", "_").replace(" ", "_")
            path = out_dir / f"{safe_name}.json"
            path.write_text(json.dumps(payload, indent=2, default=str))
            written.append(path)
    finally:
        wb_formulas.close()
        wb_values.close()
    return written
